# -*- coding: utf-8 -*-
"""
Aggiorna la colonna 'link_web' in maxi-impianti.xlsx.

Per ogni impianto scrive (o riscrive) il link cliccabile alla scheda web del
sito, nella forma:  https://www.diessemedia.it/maxi.html#imp=<CODICE>

Questi link servono a essere inviati ai clienti (aprono la pagina Maxi con la
scheda dell'impianto gia' aperta). La colonna viene rigenerata da zero a ogni
esecuzione, quindi resta sempre allineata ai codici presenti (nuovi impianti
inclusi). La colonna 'link_web' NON viene letta da build-maxi-data.py: e' solo
una comodita' per il lavoro commerciale.

USO: viene lanciato in automatico da AGGIORNA_MAPPA_MAXI.bat
     (oppure a mano:  python aggiorna-link-schede.py)
"""
import os, sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX = 'maxi-impianti.xlsx'
BASE = 'https://www.diessemedia.it/maxi.html#imp='

def main():
    if not os.path.exists(XLSX):
        print('ERRORE: non trovo %s in questa cartella.' % XLSX)
        return 1

    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Impianti'] if 'Impianti' in wb.sheetnames else wb.active

    # trova (o crea) la colonna 'link_web'
    names = {(ws.cell(row=1, column=c).value or '').strip().lower(): c
             for c in range(1, ws.max_column + 1)}
    col = names.get('link_web') or (ws.max_column + 1)
    ws.cell(row=1, column=col, value='link_web')
    ws.cell(row=2, column=col, value='Link scheda (web)')

    linkfont = Font(color='1155CC', underline='single')
    n = 0
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if code is None or str(code).strip() == '' or str(code).strip().lower() in ('code', 'codice'):
            # riga vuota o riga di etichette: svuota eventuale link residuo
            ws.cell(row=r, column=col, value=None)
            continue
        url = BASE + str(code).strip()
        cell = ws.cell(row=r, column=col, value=url)
        cell.hyperlink = url
        cell.font = linkfont
        n += 1

    ws.column_dimensions[get_column_letter(col)].width = 52

    try:
        wb.save(XLSX)
    except PermissionError:
        print('ATTENZIONE: %s e\' aperto in Excel: non ho potuto aggiornare i link.' % XLSX)
        print('            Chiudi il file e rilancia lo script (i dati del sito sono comunque a posto).')
        return 0  # non bloccante
    print('OK: colonna link_web aggiornata (%d link scritti).' % n)
    return 0

if __name__ == '__main__':
    sys.exit(main())
